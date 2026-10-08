"""CV text extraction (TXT/DOCX) + the multipart upload endpoint.

The PDF path uses pypdf, which imports `cryptography` eagerly and can't load in
some sandboxes — so here we only assert the PDF path degrades gracefully. The
TXT and DOCX paths (and the full upload→analyze flow) are exercised end to end.
"""
import io

from app.seed import load_seed
from app.services import cv_extract

CV_TEXT = b"Experienced analyst. Skills: Python, SQL, Excel. Built reporting dashboards."


def test_extract_plain_text():
    assert "Python" in cv_extract.extract_text("cv.txt", CV_TEXT)


def test_extract_docx():
    import docx

    buf = io.BytesIO()
    doc = docx.Document()
    doc.add_paragraph("Skills: Python, SQL, Excel.")
    doc.add_paragraph("Built reporting dashboards.")
    doc.save(buf)

    text = cv_extract.extract_text("cv.docx", buf.getvalue())
    assert "Python" in text and "Excel" in text


def test_extract_pdf_degrades_gracefully():
    # Not a real PDF body and no LLM configured -> empty string, never raises.
    assert cv_extract.extract_text("cv.pdf", b"%PDF-1.4 not really a pdf") == ""


def test_analyze_file_endpoint_txt(db, client):
    load_seed(db)
    r = client.post(
        "/admin/career/analyze-file",
        headers={"X-Admin-Token": "test-token"},
        files={"file": ("cv.txt", CV_TEXT, "text/plain")},
        data={"goal": "Data Analyst"},
    )
    assert r.status_code == 200
    body = r.json()
    assert {"Python", "SQL", "Excel"} <= set(body["haveSkills"])
    assert body["gapSkills"] and "roadmap" in body


def test_analyze_file_requires_admin(client):
    r = client.post(
        "/admin/career/analyze-file",
        files={"file": ("cv.txt", CV_TEXT, "text/plain")},
    )
    assert r.status_code == 401


def test_analyze_file_rejects_unextractable(db, client):
    load_seed(db)
    r = client.post(
        "/admin/career/analyze-file",
        headers={"X-Admin-Token": "test-token"},
        files={"file": ("cv.pdf", b"%PDF-bad", "application/pdf")},
    )
    assert r.status_code == 422
